"""Verify It — prove the Monthly Report Generator's numbers, with pytest.

The whole chain, end to end, in a temporary folder:

1. regenerate the sample data (seed 42) via the generator module itself,
   capturing the control totals it computed;
2. run ``generate_report`` on that folder;
3. open the .xlsx with openpyxl and assert every Summary/By-Product figure
   equals the control file — the numbers Python computed *without* Excel.

Run:
    python -m pytest verify_report.py -v
"""

import csv
import json
import sys
from pathlib import Path
#above from standard library modules
#csv imported to write csvs
#json - for converting between python objects and json
#JSON is a plain-text format for structured data
#in json keys are strings


import pytest
from openpyxl import load_workbook
#openpyxl hird-party library for reading and writing Excel

# Cross-section-style import by path (Section 15 teaches proper packaging;
# until then, sibling modules are imported the explicit way).
SECTION_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SECTION_DIR / "sample_data"))
sys.path.insert(0, str(SECTION_DIR / "project"))

import generate_sample_data 
from report_generator import generate_report  

TOLERANCE = 0.01  # one cent — floats summed in a different order may drift



#The @something syntax is a decorator — Python feature, not pytest-specific. It wraps or registers the function defined right below it
'''
@pytest.fixture
def report(...):
    ...

is shorthand for 

def report(...):
    ...
report = pytest.fixture(report)

'''
#So pytest.fixture is a function that takes your function and marks it.
#What does the mark do? It tells pytest: "this function is not a test — it's setup. 
# When a test asks for something with this name, run this function and hand over what it returns."
#scope="module" :   run once per file, share the result
#Without the scope argument, the fixture re-runs fresh for every test that uses it.

#Register report as setup. Run it once per file. 
# Any test with a report parameter gets (control, xlsx_path).
#And tmp_path_factory in its parameters shows fixtures can request other fixtures the same way

@pytest.fixture(scope="module")
def report(tmp_path_factory) -> tuple[dict, Path]:
    #we had seen tmp_path earlier which is also a fixture. tmp_path_factory is also a fixture.
    #tmp_path gives you a ready-made temporary directory as a Path object. Fresh one per test, auto-cleaned later (pytest keeps the last 3 runs, then deletes).
    #tmp_path_factory gives you a maker of temporary directories. It's not a path — it's an object with a .mktemp() method you call to create directories on demand
    #One extra step: tmp_path hands you the directory; the factory hands you the tool to make directories.
    #Why do both exist? Scope
    #tmp_path is function-scoped: a new one per test, which is exactly what makes it fresh and isolated.
    #tmp_path_factory is session-scoped — broadest possible — so any fixture at any scope can request it, and it makes directories whenever asked.
    #report fixture is module-scoped; asking it to depend on function-scoped tmp_path is contradictory
    """Generate fresh sample data + report in a temp folder, once per run."""
    data_dir = tmp_path_factory.mktemp("sales_data")
    #data_dir will be created only once per run of this file.
    control = generate_sample_data.write_sample_data(data_dir) #returns 
    xlsx_path = generate_report(data_dir, data_dir / "monthly_report.xlsx")
    return control, xlsx_path



def _sheet_rows(xlsx_path: Path, sheet_name: str) -> list[tuple]:
    #Leading underscore: convention for "internal — not a test, not for outside use."
    """Return the data rows of a sheet (header excluded), values only."""
    wb = load_workbook(xlsx_path, read_only=True)
    #Opens the .xlsx. read_only=True streams the file instead of loading everything into memory — appropriate since we only read. The catch: read-only mode keeps the file handle open until you explicitly close it.
    try:
        ws = wb[sheet_name] #wb[sheet_name] picks a sheet by name, dict-style;Wrong name raises KeyError.
        return list(ws.iter_rows(min_row=2, values_only=True))
        #iter_rows(...) walks the sheet row by row.
        #min_row=2 starts at row 2 — skips the header, so tests get only data.
        #values_only=True yields plain tuples like ("2025-01", 312, 1804, 180488.45) instead of Cell objects. That's what lets tests unpack directly: for month, orders, units, revenue in monthly_rows
        #list(...) drain the generator into a real list first as We're about to close the workbook
        #first builds the full list (workbook still open), then finally closes the workbook
    finally:
        wb.close() #finally guarantees the close runs no matter what — even if the sheet name was wrong


def test_control_file_written_and_matches_return_value(report):
    #report is a fixture defined above; (report)← fixture injection
    control, xlsx_path = report #unpacks something that was already produced. report function ran just before the first test that requested it (this one)
    control_path = xlsx_path.parent / "control_totals.json"
    assert control_path.exists()
    assert json.loads(control_path.read_text(encoding="utf-8")) == control
    #json.loads = "load from string." It takes JSON text and builds the corresponding Python object:


def test_summary_monthly_revenue_matches_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "Summary")
    monthly_rows = [row for row in rows if row[0] != "TOTAL"]

    assert len(monthly_rows) == 12, "expected one Summary row per month"
    for month, orders, units, revenue in monthly_rows:
        assert orders == control["monthly_row_counts"][month]
        assert revenue == pytest.approx(
            control["monthly_revenue"][month], abs=TOLERANCE
        ), f"revenue mismatch for {month}"
        #abs is a keyword argument to pytest.approx, setting the absolute tolerance: how far apart two numbers may be and still count as equal.


def test_summary_grand_total_matches_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "Summary")
    total_rows = [row for row in rows if row[0] == "TOTAL"]

    assert len(total_rows) == 1, "expected exactly one grand-total row"
    _, total_orders, _, total_revenue = total_rows[0]
    assert total_orders == sum(control["monthly_row_counts"].values())
    assert total_revenue == pytest.approx(
        control["grand_total_revenue"], abs=TOLERANCE
    )


def test_by_product_totals_match_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "By-Product")
    product_rows = {row[0]: row for row in rows if row[0] != "TOTAL"}

    assert set(product_rows) == set(control["product_totals"])
    #set(dict) will give a set with keys right ? 
    for product, expected in control["product_totals"].items():
        _, units, revenue, _share = product_rows[product]
        assert units == expected["units"], f"units mismatch for {product}"
        assert revenue == pytest.approx(
            expected["revenue"], abs=TOLERANCE
        ), f"revenue mismatch for {product}"


def test_mismatched_csv_is_skipped_not_fatal(tmp_path, capsys):
    #capsys : Another pytest built-in fixture — same category as tmp_path
    #capture everything the code under test prints, so assertions can examine it.
    """A CSV with the wrong columns must be warned about and ignored."""
    good = tmp_path / "sales_2025_01.csv"
    with good.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "region", "product", "units", "unit_price"])
        writer.writerow(["2025-01-15", "North", "Desk Lamp", "2", "39.99"])

    bad = tmp_path / "contacts_export.csv"
    with bad.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["name", "email", "phone"])
        writer.writerow(["Ada", "ada@example.com", "555-0100"])

    xlsx_path = generate_report(tmp_path, tmp_path / "report.xlsx")

    assert "contacts_export.csv" in capsys.readouterr().err
    '''
    captured = capsys.readouterr()
    captured.out → all stdout so far, e.g. ""
    captured.err → all stderr so far, e.g. "warning: skipping contacts_export.csv: unexpected columns ['name', 'email', 'phone']\n"
    '''
    rows = _sheet_rows(xlsx_path, "Summary")
    monthly_rows = [row for row in rows if row[0] != "TOTAL"]
    assert monthly_rows == [("2025-01", 1, 2, pytest.approx(79.98))]


def test_no_usable_rows_raises(tmp_path):
    #Does generate_report actually raise ValueError?
    """An input folder that yields nothing must fail loudly, not report emptily.

    The Refine lecture's Difference 1: the draft wrote an empty report;
    the refined generator raises instead — Section 10's Monday emailer
    needs the exception, and a confidently empty report is worse than one.
    """
    with pytest.raises(ValueError):
        #pytest.raises is a function. Calling it with an exception type returns a context manager — an object with __enter__ and __exit__ methods, which is what the with statement requires.
        #The interesting logic lives in __exit__: when the block ends, Python hands __exit__ whatever exception (if any) escaped the block
        #Got a ValueError - Report "handled, don't propagate" → block exits cleanly → test continues → passes.
        #Got nothing - Raise Failed: DID NOT RAISE <class 'ValueError'> itself.
        #Got a TypeError instead -  Decline to handle it → it propagates → test errors out.
        #so its like only for ValueError pass and that is because ValueError is passed to pytest.raises
        generate_report(tmp_path, tmp_path / "report.xlsx")
