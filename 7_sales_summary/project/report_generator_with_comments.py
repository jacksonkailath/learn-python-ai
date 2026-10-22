#!/usr/bin/env python3
"""Merge monthly sales CSVs into a formatted Excel report.

The course's project spine. This file began life as the Lecture 2 AI
draft (examples/04_first_draft_report_generator.py, kept verbatim); the
Refine lecture applied the six differences that separate the two files.

Its public interface is FROZEN — Section 10's Monday emailer and
Section 12's pandas cross-check depend on it:

    generate_report(csv_folder: Path, output_path: Path) -> Path

Scans a folder of CSV files, keeps only the ones whose header contains
the required sales columns (date, region, product, units, unit_price),
and produces a two-sheet Excel workbook: a per-month summary and a
per-product breakdown. Files that don't match are skipped with a
warning to stderr — one stray export must never poison the whole run.

Usage:
    python project/report_generator.py                    # uses ../sample_data
    python project/report_generator.py /path/to/csvs -o /path/to/report.xlsx
"""

#from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

REQUIRED_COLUMNS = {"date", "region", "product", "units", "unit_price"}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
TOTAL_FONT = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="double"))
# Excel's built-in "Bad" style: light red fill, dark red text.
LOW_REVENUE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
LOW_REVENUE_FONT = Font(color="9C0006")

COUNT_FORMAT = "#,##0"
MONEY_FORMAT = "#,##0.00"
PERCENT_FORMAT = "0.0%"


def find_csv_files(folder: Path) -> list[Path]:
    """Return all .csv files directly inside folder, sorted by name."""
    return sorted(p for p in folder.glob("*.csv") if p.is_file())


def load_sales_rows(csv_folder: Path) -> list[dict[str, Any]]:
    """Read every valid sales CSV in csv_folder and return parsed rows.

    Each file's header is checked against REQUIRED_COLUMNS before any
    rows are read — extra columns are tolerated and ignored, and header
    capitalization/stray spaces are normalized away. Files missing a
    required column are skipped with a warning to stderr, as are files
    that can't be read or decoded at all. Rows that fail to convert
    units/unit_price to numbers are skipped with a warning naming the
    file and line, instead of crashing the whole run.
    """
    rows: list[dict[str, Any]] = []
    csv_files = find_csv_files(csv_folder)

    if not csv_files:
        print(f"warning: no CSV files found in {csv_folder}", file=sys.stderr)

    for path in csv_files:
        try:
            with path.open("r", newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                #csv.DictReader(f) wraps the file object in a reader that yields each row as a dictionary keyed by column name, instead of a plain list of strings.
                #csv.reader	["2026-01-04", "North", "Widget", "3", "19.99"]
                #csv.DictReader	{"date": "2026-01-04", "region": "North", "product": "Widget", "units": "3", "unit_price": "19.99"}
                fieldnames = reader.fieldnames or []
                key_for = {name: name.strip().lower() for name in fieldnames}
                #key_for dict comprehension
                #dict comprehension → {"Date": "date", ...}

                missing = REQUIRED_COLUMNS - set(key_for.values())
                if missing:
                    print(
                        f"warning: skipping non-sales file {path.name} "
                        f"(missing column(s): {', '.join(sorted(missing))})",
                        file=sys.stderr,
                    )
                    continue

                for line_num, raw_row in enumerate(reader, start=2):
                    row = {key_for[name]: raw_row[name] for name in fieldnames}
                    #raw_row is what DictReader yields on each pass of the loop: a dict keyed by the header names exactly as they appear in the file
                    #raw_row[name] is the value from the CSV for that column, and key_for[name] is the normalized column name (lowercase, stripped). 
                    try:
                        units = int(row["units"])
                        unit_price = float(row["unit_price"])
                    except (ValueError, TypeError):
                        print(
                            f"warning: skipping row {line_num} in {path.name} "
                            f"(bad units/unit_price)",
                            file=sys.stderr,
                        )
                        continue

                    rows.append(
                        {
                            "date": row["date"],
                            "region": row["region"],
                            "product": row["product"],
                            "units": units,
                            "unit_price": unit_price,
                        }
                    )
        except (OSError, UnicodeDecodeError, csv.Error) as exc:
            #OSError alone only covers one of the three ways reading a CSV can fail
            #OSError — the filesystem said no. The file vanished between glob and open, permissions are wrong, it's on a network drive that dropped, it's a directory rather than a file. This is the one the original had.
            #UnicodeDecodeError — the bytes are there-a file on disk is just bytes, but they aren't UTF-8.Text is an interpretation of those bytes. The encoding is the rulebook for that interpretation.open(..., encoding="utf-8") commits you to one specific rulebook. When the bytes don't obey it, Python raises rather than guessing.
            #csv.Error — the bytes decoded fine, but the CSV structure is broken. The classic trigger is a quoted field containing a huge run of text with an unterminated quote.The bytes decoded to valid text — the csv module just can't parse that text as CSV.
            print(f"warning: could not read {path.name}: {exc}", file=sys.stderr)
            continue

    return rows


def summarize_by_month(rows: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    """Aggregate rows into per-month orders, units, and revenue."""
    summary: dict[str, dict[str, float]] = {}
    for row in rows:
        month = row["date"][:7]
        bucket = summary.setdefault(month, {"orders": 0, "units": 0, "revenue": 0.0})
        bucket["orders"] += 1
        bucket["units"] += row["units"]
        bucket["revenue"] += row["units"] * row["unit_price"]
    return summary


def summarize_by_product(rows: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    """Aggregate rows into per-product units and revenue."""
    summary: dict[str, dict[str, float]] = {}
    for row in rows:
        product = row["product"]
        bucket = summary.setdefault(product, {"units": 0, "revenue": 0.0})
        bucket["units"] += row["units"]
        bucket["revenue"] += row["units"] * row["unit_price"]
    return summary


def style_header_row(ws: Worksheet, num_columns: int, row: int = 1) -> None:
    """Apply bold white-on-blue formatting to the header row and freeze it."""
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"A{row + 1}"


def autosize_columns(ws: Worksheet, widths: list[int]) -> None:
    """Set sensible fixed column widths."""
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_summary_sheet(ws: Worksheet, monthly: dict[str, dict[str, float]]) -> None:
    """Write the Summary sheet: Month | Orders | Units | Revenue."""
    ws.append(["Month", "Orders", "Units", "Revenue"])
    style_header_row(ws, num_columns=4)

    months = sorted(monthly)
    revenues = [monthly[m]["revenue"] for m in months]
    avg_revenue = mean(revenues) if revenues else 0.0

    total_orders = total_units = 0
    total_revenue = 0.0

    for month in months:
        data = monthly[month]
        row_idx = ws.max_row + 1
        ws.append([month, data["orders"], data["units"], data["revenue"]])
        ws.cell(row=row_idx, column=2).number_format = COUNT_FORMAT
        ws.cell(row=row_idx, column=3).number_format = COUNT_FORMAT
        ws.cell(row=row_idx, column=4).number_format = MONEY_FORMAT

        if data["revenue"] < avg_revenue:
            for col in range(1, 5):
                cell = ws.cell(row=row_idx, column=col)
                cell.fill = LOW_REVENUE_FILL
                cell.font = LOW_REVENUE_FONT

        total_orders += data["orders"]
        total_units += data["units"]
        total_revenue += data["revenue"]

    total_row = ws.max_row + 1
    ws.append(["TOTAL", total_orders, total_units, total_revenue])
    for col in range(1, 5):
        cell = ws.cell(row=total_row, column=col)
        cell.font = TOTAL_FONT
        cell.border = TOTAL_BORDER
    ws.cell(row=total_row, column=2).number_format = COUNT_FORMAT
    ws.cell(row=total_row, column=3).number_format = COUNT_FORMAT
    ws.cell(row=total_row, column=4).number_format = MONEY_FORMAT

    autosize_columns(ws, [12, 10, 10, 16])


def write_by_product_sheet(ws: Worksheet, by_product: dict[str, dict[str, float]]) -> None:
    """Write the By-Product sheet: Product | Units | Revenue | Share of Revenue."""
    ws.append(["Product", "Units", "Revenue", "Share of Revenue"])
    style_header_row(ws, num_columns=4)

    total_revenue = sum(data["revenue"] for data in by_product.values())
    products = sorted(by_product, key=lambda p: by_product[p]["revenue"], reverse=True)

    total_units = 0
    running_revenue = 0.0

    for product in products:
        data = by_product[product]
        share = (data["revenue"] / total_revenue) if total_revenue else 0.0
        row_idx = ws.max_row + 1
        ws.append([product, data["units"], data["revenue"], share])
        ws.cell(row=row_idx, column=2).number_format = COUNT_FORMAT
        ws.cell(row=row_idx, column=3).number_format = MONEY_FORMAT
        ws.cell(row=row_idx, column=4).number_format = PERCENT_FORMAT

        total_units += data["units"]
        running_revenue += data["revenue"]

    total_row = ws.max_row + 1
    total_share = (running_revenue / total_revenue) if total_revenue else 0.0
    ws.append(["TOTAL", total_units, running_revenue, total_share])
    for col in range(1, 5):
        cell = ws.cell(row=total_row, column=col)
        cell.font = TOTAL_FONT
        cell.border = TOTAL_BORDER
    ws.cell(row=total_row, column=2).number_format = COUNT_FORMAT
    ws.cell(row=total_row, column=3).number_format = MONEY_FORMAT
    ws.cell(row=total_row, column=4).number_format = PERCENT_FORMAT

    autosize_columns(ws, [20, 10, 16, 16])


def generate_report(csv_folder: Path, output_path: Path) -> Path:
    """Merge monthly sales CSVs in csv_folder into a formatted xlsx report.

    Reads every valid sales CSV in csv_folder, summarizes the results by
    month and by product, writes a two-sheet workbook to output_path,
    and returns that path.

    Raises ValueError if the folder yields no usable sales rows — the
    empty report is the one output this program refuses to produce.
    """
    csv_folder = Path(csv_folder)
    output_path = Path(output_path)

    rows = load_sales_rows(csv_folder)
    if not rows:
        raise ValueError(
            f"No usable sales rows found in {csv_folder} — expected CSVs with "
            f"columns {sorted(REQUIRED_COLUMNS)}"
        )

    monthly = summarize_by_month(rows)
    by_product = summarize_by_product(rows)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    write_summary_sheet(summary_ws, monthly)

    product_ws = wb.create_sheet("By-Product")
    write_by_product_sheet(product_ws, by_product)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Merge monthly sales CSVs into a formatted Excel report."
    )
    parser.add_argument(
        "folder",
        nargs="?",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "sample_data",
        help="Folder containing sales CSV files (default: ../sample_data)",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "monthly_report.xlsx",
        help="Output .xlsx path (default: monthly_report.xlsx next to this script)",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    if not args.folder.is_dir():
        raise SystemExit(f"Not a folder: {args.folder}")
    try:
        result_path = generate_report(args.folder, args.output)
    except ValueError as exc:
        raise SystemExit(f"{exc}\nHint: run sample_data/generate_sample_data.py first.")
    print(f"Report written to {result_path}")
