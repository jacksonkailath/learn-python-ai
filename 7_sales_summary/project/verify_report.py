"""Verify It — prove the Monthly Report Generator's numbers, with pytest.

The whole chain, end to end, in a temporary folder:

1. regenerate the sample data (seed 42) via the generator module itself,
   capturing the control totals it computed;
2. run ``generate_report`` on that folder;
3. open the .xlsx with openpyxl and assert every Summary/By-Product figure
   equals the control file — the numbers Python computed *without* Excel.

Run:
    python -m pytest verify_report.py -v
"""

import csv
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

# Cross-section-style import by path (Section 15 teaches proper packaging;
# until then, sibling modules are imported the explicit way).
SECTION_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SECTION_DIR / "sample_data"))
sys.path.insert(0, str(SECTION_DIR / "project"))

import generate_sample_data
from report_generator import generate_report

TOLERANCE = 0.01  # one cent — floats summed in a different order may drift


@pytest.fixture(scope="module")
def report(tmp_path_factory) -> tuple[dict, Path]:
    """Generate fresh sample data + report in a temp folder, once per run."""
    data_dir = tmp_path_factory.mktemp("sales_data")
    control = generate_sample_data.write_sample_data(data_dir)
    xlsx_path = generate_report(data_dir, data_dir / "monthly_report.xlsx")
    return control, xlsx_path

'''
def report():
    ...

report = pytest.fixture(report)

'''


def _sheet_rows(xlsx_path: Path, sheet_name: str) -> list[tuple]:
    """Return the data rows of a sheet (header excluded), values only."""
    wb = load_workbook(xlsx_path, read_only=True)
    try:
        ws = wb[sheet_name]
        return list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


def test_control_file_written_and_matches_return_value(report):
    control, xlsx_path = report
    control_path = xlsx_path.parent / "control_totals.json"
    assert control_path.exists()
    assert json.loads(control_path.read_text(encoding="utf-8")) == control


def test_summary_monthly_revenue_matches_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "Summary")
    monthly_rows = [row for row in rows if row[0] != "TOTAL"]

    assert len(monthly_rows) == 12, "expected one Summary row per month"
    for month, orders, units, revenue in monthly_rows:
        assert orders == control["monthly_row_counts"][month]
        assert revenue == pytest.approx(
            control["monthly_revenue"][month], abs=TOLERANCE
        ), f"revenue mismatch for {month}"
         


def test_summary_grand_total_matches_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "Summary")
    total_rows = [row for row in rows if row[0] == "TOTAL"]

    assert len(total_rows) == 1, "expected exactly one grand-total row"
    _, total_orders, _, total_revenue = total_rows[0]
    assert total_orders == sum(control["monthly_row_counts"].values())
    assert total_revenue == pytest.approx(
        control["grand_total_revenue"], abs=TOLERANCE
    )


def test_by_product_totals_match_control(report):
    control, xlsx_path = report
    rows = _sheet_rows(xlsx_path, "By-Product")
    product_rows = {row[0]: row for row in rows if row[0] != "TOTAL"}

    assert set(product_rows) == set(control["product_totals"])
    #set(dict) = {p1,p2 ...}
    for product, expected in control["product_totals"].items():
        _, units, revenue, _share = product_rows[product]
        assert units == expected["units"], f"units mismatch for {product}"
        assert revenue == pytest.approx(
            expected["revenue"], abs=TOLERANCE
        ), f"revenue mismatch for {product}"


def test_mismatched_csv_is_skipped_not_fatal(tmp_path, capsys):
    """A CSV with the wrong columns must be warned about and ignored."""
    good = tmp_path / "sales_2025_01.csv"
    with good.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "region", "product", "units", "unit_price"])
        writer.writerow(["2025-01-15", "North", "Desk Lamp", "2", "39.99"])

    bad = tmp_path / "contacts_export.csv"
    with bad.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["name", "email", "phone"])
        writer.writerow(["Ada", "ada@example.com", "555-0100"])

    xlsx_path = generate_report(tmp_path, tmp_path / "report.xlsx")

    assert "contacts_export.csv" in capsys.readouterr().err
    '''
    captured = capsys.readouterr()
    captured.out -> stdout
    captured.err -> stderr
    '''
    rows = _sheet_rows(xlsx_path, "Summary")
    monthly_rows = [row for row in rows if row[0] != "TOTAL"]
    assert monthly_rows == [("2025-01", 1, 2, pytest.approx(79.98))]


def test_no_usable_rows_raises(tmp_path):
    """An input folder that yields nothing must fail loudly, not report emptily.

    The Refine lecture's Difference 1: the draft wrote an empty report;
    the refined generator raises instead — Section 10's Monday emailer
    needs the exception, and a confidently empty report is worse than one.
    """
    with pytest.raises(ValueError):
        generate_report(tmp_path, tmp_path / "report.xlsx")
