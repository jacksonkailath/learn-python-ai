"""Example 2 — your first Excel file from Python, in ~20 lines of openpyxl.

The minimal moves: make a Workbook, grab the active sheet, put values in
cells (two equivalent ways), append whole rows, save. No formatting yet —
that's example 3.
"""

from pathlib import Path

from openpyxl import Workbook
#openpyxl is a third-party Python library for reading and writing .xlsx files
#This line pulls the Workbook class, out of that package into your module's namespace.

OUTPUT = Path(__file__).parent / "hello_workbook.xlsx"


def main() -> None:
    wb = Workbook()          # a brand-new workbook, one empty sheet inside
    #wb is an instance of the Workbook class
    #The class is the blueprint; wb is one concrete workbook built from it.
    ws = wb.active           # that sheet — "active" = the one open on screen
    #wb.active is a property on the Workbook object.
    #it returns the worksheet currently marked as active
    ws.title = "My First Sheet"

    # Way 1: address a cell like a spreadsheet user ("A1", "B1"...)
    ws["A1"] = "Product"
    ws["B1"] = "Revenue"

    # Way 2: address a cell by row/column number — better inside loops
    ws.cell(row=2, column=1, value="USB-C Hub")
    ws.cell(row=2, column=2, value=1249.75)

    # Way 3 (the workhorse): append a whole row at once
    ws.append(["Wireless Mouse", 862.50])
    ws.append(["Desk Lamp", 431.89])

    wb.save(OUTPUT)          # nothing touches disk until this line
    print(f"Wrote {OUTPUT} — open it in Excel/LibreOffice and look around.")


if __name__ == "__main__":
    main()
