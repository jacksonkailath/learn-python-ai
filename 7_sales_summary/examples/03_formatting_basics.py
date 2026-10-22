"""Example 3 — the four formatting moves that make a sheet look professional.

1. Bold (styled) header row
2. Number formats — display formatting, the data stays numeric
3. Column widths — nobody trusts a sheet full of '#####'
4. A highlight rule — flag values below a threshold in red

Writes formatting_demo.xlsx next to this script.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).parent / "formatting_demo.xlsx"

REVENUE_TARGET = 1_000.00  # anything below this gets flagged

DEMO_ROWS = [
    ("USB-C Hub", 1_249.75),
    ("Desk Lamp", 431.89),
    ("Wireless Mouse", 862.50),
    ("27in Monitor", 5_120.00),
    ("HD Webcam", 977.35),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formatting Demo"

    # --- Move 1: a styled header row -------------------------------------
    headers = ["Product", "Revenue"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="305496")

    # --- Data + Move 2: number formats + Move 4: the highlight rule ------
    below_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")  # light red
    below_font = Font(color="9C0006")                              # dark red

    for row_num, (product, revenue) in enumerate(DEMO_ROWS, start=2):
        ws.cell(row=row_num, column=1, value=product)
        revenue_cell = ws.cell(row=row_num, column=2, value=revenue)

        # "#,##0.00" is DISPLAY formatting: the cell still holds a number,
        # so Excel formulas (and our Verify It test) keep working.
        revenue_cell.number_format = "#,##0.00" #2,987.34 ; 345.56 ; 0.56, 234.00

        if revenue < REVENUE_TARGET:
            revenue_cell.fill = below_fill
            revenue_cell.font = below_font

    # --- Move 3: column widths (measured in character units) -------------
    ws.column_dimensions[get_column_letter(1)].width = 24   # "A"
    ws.column_dimensions[get_column_letter(2)].width = 14   # "B"

    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Rows under {REVENUE_TARGET:,.2f} are highlighted in red — "
          "open it and check.")


if __name__ == "__main__":
    main()
